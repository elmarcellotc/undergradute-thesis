import docx
import pandas as pd

directory = 'outputs/tables/'

def export_to_latex(df, file_name, directory=directory):
    
    df = df.round(2)
    
    with open(directory+file_name,'w') as tf:
        tf.write(df.to_latex(index=False))


def export_to_word(df, file_name, directory=directory):
    
    df = df.round(2)
    
    doc = docx.Document()
    
    t = doc.add_table(df.shape[0]+1, df.shape[1])

    # add the header rows.
    for j in range(df.shape[-1]):
        t.cell(0,j).text = df.columns[j]

    # add the rest of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[-1]):
            t.cell(i+1,j).text = str(df.values[i,j])

    # save the doc
    doc.save(directory+file_name)